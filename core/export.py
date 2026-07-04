import io
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


NAVY = "142D64"
ORANGE = "E69500"
LIGHT_GREY = "F2F2F2"


def _style_header_row(ws, row_num: int, fill_hex: str = NAVY):
    fill = PatternFill("solid", fgColor=fill_hex)
    font_color = "FFFFFF" if fill_hex not in (LIGHT_GREY,) else "000000"
    for cell in ws[row_num]:
        cell.fill = fill
        cell.font = Font(bold=True, color=font_color)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autofit(ws, min_w: int = 12, max_w: int = 45):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(
            (len(str(c.value)) for c in col if c.value is not None),
            default=min_w,
        )
        ws.column_dimensions[letter].width = min(max(width + 2, min_w), max_w)


def _write_summary(ws, results: dict):
    fee_pct = results['pmec_fee_pct'] * 100
    period = results['period']
    total_implied = results['total_implied']

    def pct(val):
        return f"{val / total_implied * 100:.1f}%" if total_implied else "N/A"

    # ── Title ──
    ws.append([f"PMEC RECONCILIATION: PERIOD {period}", "", "", ""])
    ws.merge_cells("A1:D1")
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    ws.append([])

    # ── Returns breakdown ──
    ws.append(["PMEC RETURNS", "Employees", "Amount (ZMW)", "% of Implied Total"])
    _style_header_row(ws, ws.max_row)

    ws.append(["Successful Deductions",
               results['n_successful'],
               round(results['total_successful'], 2),
               pct(results['total_successful'])])
    ws.append(["Insufficient Funds",
               results['n_insuf'],
               round(results['total_insuf'], 2),
               pct(results['total_insuf'])])
    ws.append(["Rejected",
               results['n_rej'],
               round(results['total_rej'], 2),
               pct(results['total_rej'])])
    ws.append(["TOTAL RETURNS",
               results['n_successful'] + results['n_insuf'] + results['n_rej'],
               round(total_implied, 2),
               "100%"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=LIGHT_GREY)

    ws.append([])

    # ── Bank payment ──
    ws.append(["BANK PAYMENT", "", "", ""])
    _style_header_row(ws, ws.max_row, ORANGE)

    ws.append(["Total Successful Deductions (ZMW)", "",
               round(results['total_successful'], 2), ""])
    ws.append([f"Less: PMEC Fee ({fee_pct:.1f}%)", "",
               round(-results['pmec_fee'], 2), ""])
    ws.append(["Expected Bank Transfer (ZMW)", "",
               round(results['expected_transfer'], 2), ""])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    if results['actual_transfer'] is not None:
        ws.append(["Actual Bank Transfer (ZMW)", "",
                   round(results['actual_transfer'], 2), ""])
        discrepancy = results['transfer_discrepancy']
        ws.append(["Discrepancy", "",
                   round(discrepancy, 2),
                   "PASS" if abs(discrepancy) < 1 else "CHECK"])
        ws.append([f"Implied Fee %", "",
                   f"{results['implied_fee_pct'] * 100:.2f}%", ""])

    ws.append([])

    # ── Data quality flags ──
    ws.append(["DATA QUALITY FLAGS", "", "", ""])
    _style_header_row(ws, ws.max_row, ORANGE)

    flags = []
    if results['n_insuf_missing_amount'] > 0:
        flags.append(
            f"{results['n_insuf_missing_amount']} insufficient funds records "
            "have no amount recorded"
        )
    if results['n_rej_duplicate_employees'] > 0:
        flags.append(
            f"{results['n_rej_duplicate_employees']} rejection rows share a "
            "duplicate employee number"
        )
    if not flags:
        flags.append("No data quality issues detected")
    for flag in flags:
        ws.append([flag, "", "", ""])

    _autofit(ws)


def build_excel(
    df_success: pd.DataFrame,
    df_insuf: pd.DataFrame,
    df_rej: pd.DataFrame,
    results: dict,
) -> bytes:
    """Build the reconciliation workbook in memory and return raw bytes.

    Sheets:
        1_Reconciliation_Summary   — formatted summary with PMEC fee and flags
        2_Successful_Payments      — full successful deductions table
        3_Insufficient_Funds       — full insufficient funds table
        4_Rejections               — full rejections table
        5a_Success_Breakdown       — successful payments by Personnel Area
        5b_Insuf_Breakdown         — insufficient funds by Reason
        5c_Rejection_Breakdown     — rejections by category (COMMENTS)
    """
    rej_cols = [
        "Employee No", "Name", "Start Date", "End Date",
        "Amount (ZMW)", "Rejection Reason", "Status",
    ]

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # Write data sheets first so openpyxl book is initialised
        df_success.to_excel(writer, sheet_name="2_Successful_Payments", index=False)
        df_insuf.to_excel(writer, sheet_name="3_Insufficient_Funds", index=False)
        df_rej[rej_cols].to_excel(writer, sheet_name="4_Rejections", index=False)
        results['success_breakdown'].to_excel(
            writer, sheet_name="5a_Success_Breakdown", index=False)
        results['insuf_breakdown'].to_excel(
            writer, sheet_name="5b_Insuf_Breakdown", index=False)
        results['rej_breakdown'].to_excel(
            writer, sheet_name="5c_Rejection_Breakdown", index=False)

        # Build summary sheet on the live workbook object, insert at position 0
        wb = writer.book
        ws_summary = wb.create_sheet("1_Reconciliation_Summary", 0)
        _write_summary(ws_summary, results)

        # Style header row of every data sheet
        for name in [
            "2_Successful_Payments", "3_Insufficient_Funds", "4_Rejections",
            "5a_Success_Breakdown", "5b_Insuf_Breakdown", "5c_Rejection_Breakdown",
        ]:
            ws = wb[name]
            _style_header_row(ws, 1)
            _autofit(ws)

    return buffer.getvalue()
